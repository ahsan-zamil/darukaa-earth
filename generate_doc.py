import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Styles
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(100, 116, 139)
    p.paragraph_format.space_after = Pt(20)
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_paragraph(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(51, 65, 85)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    p.paragraph_format.space_after = Pt(3)
    return p

# --- DOCUMENT CONTENT GENERATION ---

add_title("Darukaa.Earth Platform")
add_subtitle("Carbon & Biodiversity Project Intelligence Platform — Hackathon Submission Document")

# Section 1: Submission Links & Details
add_heading_1("1. Submission Links & Key Credentials")

add_paragraph("[Insert Your GitHub Repository URL Here - e.g. https://github.com/your-username/darukaa-earth]", "GitHub Repository Link: ")
add_paragraph("https://fresh-adults-read.loca.lt", "Live Demo URL: ")
add_paragraph("demo@darukaa.earth", "Administrator Demo Email: ")
add_paragraph("Demo@12345", "Administrator Demo Password: ")

# Section 2: Repository Access Instructions
add_heading_1("2. Repository Access Instructions")
add_paragraph("Access has been granted to the following reviewer email accounts (both @darukaa.com and @darukaa.earth domains):")
add_bullet("ankita.dasgupta@darukaa.com / ankita.dasgupta@darukaa.earth")
add_bullet("harsh.kumar@darukaa.com / harsh.kumar@darukaa.earth")
add_bullet("utkarsh.gauniyal@darukaa.com / utkarsh.gauniyal@darukaa.earth")
add_bullet("guneet.mutreja@darukaa.com / guneet.mutreja@darukaa.earth")

# Section 3: Architecture & System Overview
add_heading_1("3. System Architecture & Tech Stack")
add_paragraph("Darukaa.Earth is built as a full-stack, enterprise-grade geospatial analytics platform with a decoupled single-page application (SPA) frontend and a high-performance GIS backend REST service.")

add_heading_2("Tech Stack Summary")
add_bullet("React 18, TypeScript, Vite, Tailwind CSS, Lucide Icons", "Frontend Layer: ")
add_bullet("Mapbox GL JS & @mapbox/mapbox-gl-draw (Interactive vector polygon rendering & drawing)", "Mapping Engine: ")
add_bullet("Chart.js & react-chartjs-2 (12-month ecological time-series trends)", "Visualization: ")
add_bullet("Python 3.12 with FastAPI, Pydantic v2, PyJWT, Passlib (bcrypt)", "Backend Framework: ")
add_bullet("PostgreSQL 16 with PostGIS 3.4 Extension (Spatial POLYGON geometries & SRID 4326)", "Database & GIS: ")
add_bullet("SQLAlchemy 2.0, GeoAlchemy2, Alembic", "ORM & Migrations: ")
add_bullet("Docker Compose, GitHub Actions CI/CD pipeline, pre-commit hooks (Ruff, Black, ESLint)", "DevOps & Quality: ")

# Section 4: Database Schema Breakdown
add_heading_1("4. Database Schema Breakdown")
add_paragraph("The database leverages PostgreSQL with the native PostGIS spatial extension enabled. The key entities include:")

add_bullet("Stores administrator user profiles, role-based metadata, and bcrypt password hashes.", "USERS: ")
add_bullet("Stores project records, project types (Carbon, Biodiversity, Mixed), project statuses, and country locations.", "PROJECTS: ")
add_bullet("Stores geographical site boundaries using PostGIS POLYGON (SRID 4326) and calculated spatial area in hectares (computed via ST_Area).", "SITES: ")
add_bullet("Stores historical 12-month ecological metric records (Carbon Stock tCO2e/ha, Soil Organic Carbon %, Species Richness Index, Soil pH, Soil Moisture %, Rainfall mm, Deforestation Index).", "SITE_METRICS: ")

# Section 5: Local Setup & Running Instructions
add_heading_1("5. Quickstart & Local Setup Instructions")

add_heading_2("Option 1: Docker Compose (Recommended)")
add_paragraph("1. Clone the repository and copy environment file:")
add_paragraph("git clone https://github.com/your-username/darukaa-earth.git\ncd darukaa-earth\ncp .env.example .env", bold_prefix="   Commands: ")
add_paragraph("2. Build and launch all containers:")
add_paragraph("docker-compose up --build", bold_prefix="   Commands: ")
add_paragraph("3. Access running services:")
add_bullet("http://localhost:5173", "Frontend SPA: ")
add_bullet("http://localhost:8000/api/docs", "Backend Swagger API Docs: ")

add_heading_2("Option 2: Manual Local Setup")
add_bullet("Run backend: navigate to backend/, pip install -r requirements.txt, alembic upgrade head, python seed.py, uvicorn app.main:app --reload --port 8000", "Backend: ")
add_bullet("Run frontend: navigate to frontend/, npm install, npm run dev", "Frontend: ")

# Section 6: CI/CD Pipeline & Quality Enforcement
add_heading_1("6. CI/CD Pipeline & Code Quality")
add_paragraph("Automated GitHub Actions CI workflow is configured in .github/workflows/ci.yml and triggers on every push and pull request:")
add_bullet("Spawns a PostGIS 16 service container, installs dependencies, executes Ruff & Black linter checks, applies Alembic migrations, and runs full Pytest test suite (100% pass).", "Backend CI Job: ")
add_bullet("Sets up Node.js 20, installs npm dependencies, executes ESLint zero-warning check, performs strict TypeScript type-checking, and verifies Vite production build.", "Frontend CI Job: ")
add_bullet("Configured via .pre-commit-config.yaml to enforce code standards before git commit.", "Pre-commit Hooks: ")

# Section 7: Notes & Engineering Trade-offs
add_heading_1("7. Implementation Notes & Trade-offs")
add_bullet("Synthetic environmental metric datasets were generated via seed script to allow immediate high-fidelity Chart.js visualization without requiring paid satellite API keys.", "Synthetic Ecological Data: ")
add_bullet("All site polygon surface areas are computed in geodesic space (WGS84) directly using PostGIS ST_Area(geometry::geography) / 10000.0 for accurate hectare calculations.", "PostGIS Spatial Engine: ")
add_bullet("If a Mapbox API token is missing, the application renders an inline warning banner while maintaining 100% functionality of all dashboards, tables, and CRUD operations.", "Mapbox Fallback Handling: ")

doc.save("c:/Users/ASUS/Desktop/darukaa-earth/Darukaa_Earth_Submission_Document.docx")
print("Submission document successfully created!")
